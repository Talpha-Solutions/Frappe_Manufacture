# Copyright (c) 2026, talpha solutions and contributors
# For license information, please see license.txt

"""Excel export of a Tender: a summary sheet across its tender manifests, then one bill per manifest.

Detailed and simplified follow the Tender Manifest export: detailed is the internal costing (cost,
margin, specs, notes); simplified is the client's copy at tender price, item lines only.
"""

import re
from io import BytesIO

import frappe
from frappe import _
from frappe.utils import flt, formatdate
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Side
from openpyxl.utils import get_column_letter

from fitzgerald_kitchens.fitzgerald_kitchens.doctype.tender_manifest.tender_manifest_excel import (
	FONT,
	HEADER_FILL,
	MONEY,
	QTY,
	THIN,
	TOTAL_FILL,
	_box_row,
	_header_lines,
	_page_setup,
	_write_bill,
	_write_title,
)

# Summary columns: (heading, manifest field or key, kind); kind sets the format and alignment
DETAILED_COLUMNS = [
	("#", "number", "ref"),
	("Unit Type", "tender_name", "text"),
	("Tender Manifest", "name", "text"),
	("Units", "kitchens_to_tender", "qty"),
	("Cost / Unit", "total_cost_per_kitchen", "money"),
	("Cost Base", "cost_base", "money"),
	("Margin %", "target_margin_pct", "pct"),
	("Margin", "margin_amount", "money"),
	("Price / Unit", "tender_price_per_kitchen", "money"),
	("Tender Price", "tender_price_total", "money"),
]
SIMPLIFIED_COLUMNS = [
	("#", "number", "ref"),
	("Unit Type", "tender_name", "text"),
	("Units", "kitchens_to_tender", "qty"),
	("Price / Unit", "tender_price_per_kitchen", "money"),
	("Tender Price", "tender_price_total", "money"),
]
# columns added up in the total row (per-unit figures and margin % don't add up)
SUMMED = {"kitchens_to_tender", "cost_base", "margin_amount", "tender_price_total"}
WIDTHS = {"ref": 5, "text": 34, "qty": 9, "money": 15, "pct": 10}
FIRST_COL = 2


@frappe.whitelist()
def download_tender_excel(name: str, simplified: int | str = 0):
	tender = frappe.get_doc("Tender", name)
	tender.check_permission("read")
	client = bool(int(simplified or 0))

	manifest_names = frappe.get_all(
		"Tender Manifest", filters={"tender": name}, pluck="name", order_by="creation asc"
	)
	if not manifest_names:
		frappe.throw(_("Add a tender manifest to this tender before downloading it."))
	manifests = [frappe.get_doc("Tender Manifest", m) for m in manifest_names]
	for manifest in manifests:
		manifest.check_permission("read")

	wb = Workbook()
	summary = wb.active
	summary.title = "Summary"

	# one bill per manifest, each on its own sheet
	sheet_names = []
	for number, manifest in enumerate(manifests, start=1):
		sheet_name = _sheet_name(number, manifest.tender_name or manifest.name, sheet_names)
		sheet_names.append(sheet_name)
		_write_bill(
			wb.create_sheet(sheet_name),
			manifest,
			_header_lines(manifest, client),
			client,
			title=manifest.tender_name or manifest.name,
			with_totals=True,
		)

	_write_tender_summary(summary, tender, manifests, sheet_names, client)

	out = BytesIO()
	wb.save(out)
	suffix = " (Simplified)" if client else ""
	frappe.response["filename"] = f"{tender.name} - {tender.tender_name}{suffix}.xlsx".replace("/", "-")
	frappe.response["filecontent"] = out.getvalue()
	frappe.response["type"] = "binary"


def _write_tender_summary(ws, tender, manifests, sheet_names, client: bool):
	columns = SIMPLIFIED_COLUMNS if client else DETAILED_COLUMNS
	last_col = FIRST_COL + len(columns) - 1
	ws.column_dimensions["A"].width = 2
	for col, (_heading, _key, kind) in enumerate(columns, start=FIRST_COL):
		ws.column_dimensions[get_column_letter(col)].width = WIDTHS[kind]

	row = _write_title(ws, "Tender Summary", _tender_header(tender, client))

	header_row = row
	for col, (heading, _key, kind) in enumerate(columns, start=FIRST_COL):
		cell = ws.cell(row=row, column=col, value=heading)
		cell.font = Font(name=FONT, size=10, bold=True)
		cell.fill = HEADER_FILL
		cell.alignment = Alignment(horizontal="left" if kind in ("ref", "text") else "right", vertical="center")
	_box_row(ws, row, FIRST_COL, last_col, top=THIN, bottom=THIN)
	ws.freeze_panes = ws.cell(row=row + 1, column=1)
	row += 1

	for number, (manifest, sheet_name) in enumerate(zip(manifests, sheet_names), start=1):
		for col, (_heading, key, kind) in enumerate(columns, start=FIRST_COL):
			value = number if key == "number" else manifest.get(key)
			cell = ws.cell(row=row, column=col, value=_value(value, kind))
			_style(cell, kind)
			if key == "tender_name":
				# jump to the manifest's sheet
				cell.value = value or manifest.name
				cell.hyperlink = f"#'{sheet_name}'!A1"
				cell.font = Font(name=FONT, size=10, color="1F4E99", underline="single")
		_box_row(ws, row, FIRST_COL, last_col)
		row += 1

	# total row: units, cost base, margin and tender price add up across the manifests
	ws.cell(row=row, column=FIRST_COL + 1, value=_("Total"))
	for col, (_heading, key, kind) in enumerate(columns, start=FIRST_COL):
		if key in SUMMED:
			cell = ws.cell(row=row, column=col, value=sum(flt(m.get(key)) for m in manifests))
			_style(cell, kind)
		if key == "target_margin_pct":
			# overall margin on cost, since manifests can have different margins
			cost = sum(flt(m.cost_base) for m in manifests)
			margin = sum(flt(m.margin_amount) for m in manifests)
			cell = ws.cell(row=row, column=col, value=margin / cost * 100 if cost else 0)
			_style(cell, kind)
	for col in range(FIRST_COL, last_col + 1):
		cell = ws.cell(row=row, column=col)
		cell.fill = TOTAL_FILL
		cell.font = Font(name=FONT, size=10, bold=True)
	_box_row(ws, row, FIRST_COL, last_col, top=THIN, bottom=Side(style="double", color="808080"))

	_page_setup(ws, header_row, last_column=get_column_letter(last_col))


def _tender_header(tender, client: bool) -> list[str]:
	company = frappe.defaults.get_user_default("Company") or frappe.db.get_single_value(
		"Global Defaults", "default_company"
	)
	dates = " · ".join(
		filter(
			None,
			[
				f"Tender date: {formatdate(tender.tender_date)}" if tender.tender_date else "",
				f"Submission due: {formatdate(tender.submission_due_date)}" if tender.submission_due_date else "",
			],
		)
	)
	site = ", ".join(line.strip() for line in (tender.site_address or "").splitlines() if line.strip())
	return [
		f"From: {company}" if company else "",
		f"Tender: {tender.tender_name} ({tender.name})",
		f"Client: {tender.party_title}" if tender.party_name else "",
		f"Site: {site}" if site else "",
		dates,
		"" if client else f"Status: {tender.status}",
	]


def _value(value, kind: str):
	if kind in ("money", "qty", "pct"):
		return flt(value)
	return value


def _style(cell, kind: str):
	cell.font = Font(name=FONT, size=10)
	cell.alignment = Alignment(horizontal="left" if kind in ("ref", "text") else "right", vertical="top")
	if kind == "money":
		cell.number_format = MONEY
	elif kind == "qty":
		cell.number_format = QTY
	elif kind == "pct":
		cell.number_format = '0.00"%"'


def _sheet_name(number: int, title: str, taken: list[str]) -> str:
	"""Excel sheet names: at most 31 characters, none of []:*?/\\ and unique in the workbook."""
	base = re.sub(r"[\[\]:*?/\\]", "-", f"{number}. {title}").strip()[:31].rstrip()
	name, n = base, 2
	while name in taken or name == "Summary":
		suffix = f" ({n})"
		name, n = base[: 31 - len(suffix)] + suffix, n + 1
	return name
